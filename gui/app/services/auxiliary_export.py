"""辅助查询快照导出；与界面共用已查询的记录，不重新读库。"""

from __future__ import annotations

import csv
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DAILY_COLUMNS = [
    ("shop", "店铺"),
    ("负向申诉", "负向申诉提交数"),
    ("复议", "复议提交数"),
    ("售后申诉", "售后申诉提交数"),
    ("total", "合计"),
]
SUMMARY_COLUMNS = [
    ("shop", "店铺"),
    ("business_type", "申诉类型"),
    ("submitted", "范围内提交数"),
    ("pass_rate", "审核通过率"),
    ("submission_pass_rate", "通过占提交比例"),
    ("passed", "计入通过数"),
    ("reviewed", "已出审核结果"),
    ("approved", "完全通过"),
    ("partial", "部分通过"),
    ("rejected", "未通过"),
    ("waiting", "待审核"),
    ("unknown", "结果待确认"),
    ("withdrawn", "已撤回"),
]
DETAIL_COLUMNS = [
    ("business_type", "申诉类型"),
    ("shop", "店铺"),
    ("order_no", "订单号"),
    ("outcome", "当前审核结果"),
    ("submission_time", "统计提交时间（北京时间）"),
    ("appeal_item", "问题类型/申诉项"),
    ("record_key", "申诉记录标识"),
    ("aftersale_id", "售后编号"),
    ("goods_title", "商品名称"),
    ("time_source", "提交时间来源"),
    ("local_submission_time", "程序提交时间（北京时间）"),
    ("platform_submission_time", "平台提交时间（北京时间）"),
    ("counts_as_pass", "计入通过数"),
    ("counts_as_reviewed", "计入已出审核数"),
    ("last_checked", "最后回查时间（北京时间）"),
    ("result_text", "结果原文/说明"),
    ("amount", "申诉/涉及金额"),
    ("reason", "申请原因"),
    ("description", "申请说明"),
]


def display_value(row: dict, key: str):
    value = row.get(key)
    if isinstance(value, bool):
        return "是" if value else "否"
    if key in {"pass_rate", "submission_pass_rate"}:
        return f"{value:.1%}" if value is not None else "—"
    return "" if value is None else value


def _sheet(workbook, title, columns, rows):
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    for index, (key, _label) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (
            25 if "time" in key or key == "last_checked" else 19
        )
    if title == "查询说明":
        sheet.column_dimensions["A"].width = 35
        sheet.column_dimensions["B"].width = 110
    header = []
    for _key, label in columns:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        header.append(cell)
    sheet.append(header)
    for row in rows:
        cells = []
        for key, _label in columns:
            value = row.get(key)
            if isinstance(value, bool):
                value = "是" if value else "否"
            if value is None:
                value = "—" if key in {"pass_rate", "submission_pass_rate"} else ""
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                # Order IDs and text stay literal, including strings beginning with '='.
                cell.data_type = "s"
                if title == "查询说明":
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            if key in {"pass_rate", "submission_pass_rate"} and isinstance(
                value, (int, float)
            ):
                cell.number_format = "0.0%"
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def _metadata(report: dict) -> list[dict]:
    criteria = report["criteria"]
    return [
        {"item": key, "value": value}
        for key, value in (
            ("查询开始日期（北京时间）", criteria["start"]),
            ("查询结束日期（包含当天）", criteria["end"]),
            ("筛选店铺", criteria["shop"] or "全部店铺"),
            ("筛选申诉类型", criteria["business_type"] or "全部类型"),
            ("售后部分通过计入通过数", "是" if criteria["partial_pass"] else "否"),
            ("提交数量统计范围", f"{criteria['start']} 至 {criteria['end']}（包含结束当天）"),
            ("数据快照时间（北京时间）", report["generated_at"]),
            ("明细记录数", len(report["details"])),
            ("缺少提交时间的记录数（未计入）", report["missing_time"]),
            ("统计说明", report["policy"]),
            ("审核通过率", "计入通过数 ÷ 已出审核结果数"),
            ("通过占提交比例", "计入通过数 ÷ 查询范围内提交数"),
        )
    ]


def export_report(report: dict, out_path: str | Path, kind: str = "xlsx") -> None:
    if kind not in {"xlsx", "csv"}:
        raise ValueError("不支持的导出格式。")
    path = Path(out_path)
    with tempfile.NamedTemporaryFile(
        dir=path.parent, prefix=".appeal-export-", suffix="." + kind, delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        if kind == "xlsx":
            workbook = Workbook(write_only=True)
            _sheet(
                workbook,
                "查询说明",
                [("item", "项目"), ("value", "说明")],
                _metadata(report),
            )
            _sheet(workbook, "提交数量", DAILY_COLUMNS, report["daily"])
            _sheet(workbook, "通过率汇总", SUMMARY_COLUMNS, report["summary"])
            _sheet(workbook, "订单结果明细", DETAIL_COLUMNS, report["details"])
            workbook.save(temporary)
            workbook.close()
        else:
            with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                columns = [
                    ("query_start", "查询开始日期"),
                    ("query_end", "查询结束日期"),
                    *DETAIL_COLUMNS,
                ]
                writer.writerow([label for _key, label in columns])
                for detail in report["details"]:
                    row = {
                        **detail,
                        "query_start": report["criteria"]["start"],
                        "query_end": report["criteria"]["end"],
                    }
                    values = []
                    for key, _label in columns:
                        value = display_value(row, key)
                        if isinstance(value, str) and value.lstrip().startswith(
                            ("=", "+", "-", "@")
                        ):
                            value = "'" + value
                        values.append(value)
                    writer.writerow(values)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
