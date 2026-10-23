"""Read account credentials from an Excel workbook without retaining secrets."""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path

_HEADER_ALIASES = {
    "username": {"账号", "帐号", "账户", "登录账号", "账号名", "子账号", "username", "account"},
    "password": {"密码", "登录密码", "账号密码", "password"},
    "shop": {"店铺", "店铺名", "店铺名称", "店名", "shop", "shop_name"},
    "group": {"分组", "所属组", "店铺分组", "group", "group_name"},
}
_INVALID_PATH_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_RESERVED_NAMES = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}


@dataclass
class AccountRecord:
    username: str
    password: str
    shop: str = ""
    group_name: str = ""

    def clear_password(self) -> None:
        self.password = ""


def _header_mapping(values) -> dict[str, int | None] | None:
    normalized = [str(value or "").strip().lower().rstrip(":：") for value in values]
    mapping = {
        field: next((index for index, value in enumerate(normalized) if value in aliases), None)
        for field, aliases in _HEADER_ALIASES.items()
    }
    return mapping if mapping["username"] is not None and mapping["password"] is not None else None


def parse_pasted_accounts(text: str) -> list[AccountRecord]:
    """Parse Excel/tab or CSV rows in account,password,shop,group order."""
    text = text.lstrip("\ufeff")
    raw_rows = list(csv.reader(io.StringIO(text), delimiter="\t" if "\t" in text else ","))
    raw_rows = [row for row in raw_rows if any(cell != "" for cell in row)]
    if not raw_rows:
        return []
    mapping = _header_mapping(raw_rows[0])
    if mapping is not None:
        records = []
        for row in raw_rows[1:]:
            def value(field, current_row=row):
                index = mapping[field]
                return current_row[index] if index is not None and index < len(current_row) else ""

            records.append(
                AccountRecord(value("username"), value("password"), value("shop"), value("group"))
            )
        return records
    if any(len(row) > 4 for row in raw_rows):
        raise ValueError("最多粘贴 4 列：账号、密码、店铺名、分组。")
    return [AccountRecord(*(row + [""] * (4 - len(row)))) for row in raw_rows]


def _text_cell(cell, *, sheet: str, row: int, field: str) -> str:
    value = cell.value if cell is not None else None
    if cell is not None and cell.data_type == "f":
        raise ValueError(f"工作表「{sheet}」第 {row} 行的账号或密码含公式，请先转为文本值。")
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        raise ValueError(f"工作表「{sheet}」第 {row} 行的{field}不是文本或整数。")
    text = str(int(value))
    if len(text.lstrip("-")) > 15:
        raise ValueError(f"工作表「{sheet}」第 {row} 行的数字{field}过长，请改成文本后核对精度。")
    number_format = getattr(cell, "number_format", "")
    return text.zfill(len(number_format)) if number_format and set(number_format) == {"0"} else text


def _safe_profile_name(value: str) -> str:
    value = _INVALID_PATH_CHARS.sub("_", value).strip().rstrip(". ")
    # Must match paths._slugify_shop(), otherwise the profile directory found by
    # the scheduler and the account-list key would appear as two different shops.
    value = re.sub(r"[^\w-]", "_", value).strip("_") or "店铺"
    if value.upper() in _RESERVED_NAMES:
        value = f"店铺_{value}"
    if len(value) > 100:
        digest = sha256(value.encode("utf-8")).hexdigest()[:8]
        value = f"{value[:88]}_{digest}"
    return value


def assign_account_shop_names(records: list[AccountRecord]) -> list[AccountRecord]:
    """Validate, deduplicate by account and assign safe unique profile names."""
    deduplicated: dict[str, AccountRecord] = {}
    for number, record in enumerate(records, 1):
        username = str(record.username or "").strip()
        password = str(record.password or "")
        shop = str(record.shop or "").strip()
        group = str(record.group_name or "").strip()
        if not username or not password:
            raise ValueError(f"第 {number} 行的账号和密码均不能为空。")
        if len(username) > 200 or len(password) > 500 or len(shop) > 120 or len(group) > 120:
            raise ValueError(f"第 {number} 行的账号、密码、店铺名或分组过长。")
        previous = deduplicated.get(username)
        if previous is not None:
            shop = shop or previous.shop
            group = group or previous.group_name
        deduplicated[username] = AccountRecord(username, password, shop, group)
    if not deduplicated:
        raise ValueError("没有可保存的账号。")
    records = list(deduplicated.values())
    bases = [record.shop or re.split(r"[:：]", record.username, maxsplit=1)[0].strip() or record.username
             for record in records]
    counts = Counter(bases)
    used: set[str] = set()
    result: list[AccountRecord] = []
    for record, base in zip(records, bases):
        candidate = base
        if not record.shop and counts[base] > 1:
            parts = re.split(r"[:：]", record.username, maxsplit=1)
            suffix = parts[1].strip() if len(parts) > 1 else sha256(record.username.encode()).hexdigest()[:6]
            candidate = f"{base}（账号{suffix}）"
        candidate = _safe_profile_name(candidate)
        if candidate in used:
            candidate = _safe_profile_name(f"{candidate}_{sha256(record.username.encode()).hexdigest()[:6]}")
        used.add(candidate)
        result.append(AccountRecord(record.username, record.password, candidate, record.group_name))
    return result


def read_excel_rows(
    path: str | Path,
    sheet_name: str,
    start_row: int,
    columns: dict[str, int | None],
) -> list[AccountRecord]:
    """Read explicitly selected zero-based columns from one worksheet."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    records = []
    try:
        sheet = workbook[sheet_name]
        for row_number, cells in enumerate(sheet.iter_rows(min_row=start_row), start_row):
            def cell(field, current_cells=cells):
                index = columns.get(field)
                return current_cells[index] if index is not None and index < len(current_cells) else None

            username = _text_cell(cell("username"), sheet=sheet.title, row=row_number, field="账号").strip()
            password = _text_cell(cell("password"), sheet=sheet.title, row=row_number, field="密码")
            shop = str(cell("shop").value or "").strip() if cell("shop") is not None else ""
            group = str(cell("group").value or "").strip() if cell("group") is not None else ""
            if not username and not password and not shop and not group:
                continue
            records.append(AccountRecord(username, password, shop, group))
    finally:
        workbook.close()
    return assign_account_shop_names(records)


def inspect_excel_layout(path: str | Path) -> dict[str, dict]:
    """Return sheet sizes and the first recognized header within 20 rows."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    result = {}
    try:
        for sheet in workbook.worksheets:
            header_row = 1
            mapping = {field: None for field in _HEADER_ALIASES}
            for row_number, cells in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 20)), 1):
                detected = _header_mapping([cell.value for cell in cells])
                if detected is not None:
                    header_row = row_number
                    mapping = detected
                    break
            result[sheet.title] = {
                "max_column": max(1, sheet.max_column),
                "header_row": header_row,
                "mapping": mapping,
            }
    finally:
        workbook.close()
    return result


def read_excel_accounts(path: str | Path) -> list[AccountRecord]:
    """Find an account/password header in the first 20 rows of any sheet."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception:  # noqa: BLE001 - normalize workbook/parser failures for operators
        raise ValueError("无法打开 Excel，请选择未加密的 .xlsx 文件。") from None
    rows: list[AccountRecord] = []
    try:
        for sheet in workbook.worksheets:
            mapping = None
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                if mapping is None:
                    mapping = _header_mapping([cell.value for cell in cells])
                    if mapping is not None:
                        continue
                    if row_number >= 20:
                        break
                    continue
                username_cell = cells[mapping["username"]] if mapping["username"] < len(cells) else None
                password_cell = cells[mapping["password"]] if mapping["password"] < len(cells) else None
                shop_index = mapping["shop"]
                group_index = mapping["group"]
                shop_cell = cells[shop_index] if shop_index is not None and shop_index < len(cells) else None
                group_cell = cells[group_index] if group_index is not None and group_index < len(cells) else None
                username = _text_cell(username_cell, sheet=sheet.title, row=row_number, field="账号").strip()
                password = _text_cell(password_cell, sheet=sheet.title, row=row_number, field="密码")
                explicit_shop = str(shop_cell.value or "").strip() if shop_cell is not None else ""
                group_name = str(group_cell.value or "").strip() if group_cell is not None else ""
                if not username and not password:
                    continue
                if not username or not password:
                    raise ValueError(f"工作表「{sheet.title}」第 {row_number} 行的账号和密码必须同时填写。")
                if len(username) > 200 or len(password) > 500 or len(explicit_shop) > 120:
                    raise ValueError(f"工作表「{sheet.title}」第 {row_number} 行的账号、密码或店铺名过长。")
                rows.append(AccountRecord(username, password, explicit_shop, group_name))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("未找到“账号/子账号”和“密码”两列，或表格中没有有效账号。")

    return assign_account_shop_names(rows)
