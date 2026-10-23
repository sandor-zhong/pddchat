from __future__ import annotations

import os
import sys
import threading
import time
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "tests"))
from PyQt5.QtCore import Qt
from PyQt5.QtTest import QTest
from PyQt5.QtWidgets import QApplication, QPushButton, QWidget
from test_auxiliary_queries import NOW, _DatabaseFixture

from gui.app import main_window
from gui.app.pages.auxiliary_queries_page import AuxiliaryQueriesPage, QueryTableModel
from gui.app.services import auxiliary_queries as queries


class AuxiliaryPageTests(_DatabaseFixture):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        super().setUp()
        self.enterContext(patch.object(queries, "beijing_now", return_value=NOW))
        self.insert("orders", submitted_at="2026-10-14T00:00:00Z")
        self.insert(
            "aftersale_appeal_records",
            shop="店铺B",
            status="部分通过",
            submitted_at="2026-10-14T01:00:00Z",
        )
        self.page = AuxiliaryQueriesPage()
        self.addCleanup(self.page.close)

    def report(self, **kwargs):
        return queries.query_report(
            date(2026, 10, 1), date(2026, 10, 14), now=NOW, database=self.path, **kwargs
        )

    def wait_for(self, predicate):
        deadline = time.monotonic() + 5
        while not predicate() and time.monotonic() < deadline:
            self.app.processEvents()
            QTest.qWait(10)
        self.assertTrue(predicate())

    def show_report(self):
        self.page._on_result(self.page._criteria(), self.report())

    def test_query_populates_all_three_tables_and_export_uses_exact_snapshot(self):
        report = self.report()
        with (
            patch.object(queries, "query_report", return_value=report),
            patch("gui.app.pages.auxiliary_queries_page.msgbox.information"),
            patch(
                "gui.app.pages.auxiliary_queries_page.QFileDialog.getSaveFileName",
                return_value=(str(self.root / "明细.xlsx"), ""),
            ),
        ):
            self.page.query()
            self.wait_for(lambda: self.page._report is not None)
            self.assertEqual(self.page._daily_model.rowCount(), 2)
            self.assertEqual(self.page._detail_model.rowCount(), 2)
            self.assertTrue(self.page._excel_button.isEnabled())
            with patch.object(
                queries,
                "query_report",
                side_effect=AssertionError("export must not requery"),
            ):
                self.page._export("xlsx")
                self.wait_for(lambda: not self.page.is_exporting())
            self.assertTrue((self.root / "明细.xlsx").is_file())
            from openpyxl import load_workbook

            workbook = load_workbook(self.root / "明细.xlsx")
            try:
                self.assertEqual(workbook["订单结果明细"].max_row - 1, 2)
            finally:
                workbook.close()

    def test_changed_filters_disable_old_export_and_keep_existing_results_visible(self):
        self.show_report()
        self.assertTrue(self.page._csv_button.isEnabled())
        self.page._shop_combo.setEditText("店铺B")
        self.assertFalse(self.page._csv_button.isEnabled())
        self.assertEqual(self.page._detail_model.rowCount(), 2)
        self.assertIn("筛选条件已修改", self.page._snapshot_label.text())

    def test_filter_change_during_async_query_discards_old_result(self):
        report = self.report()
        started, release = threading.Event(), threading.Event()

        def wait_query(*_args, **_kwargs):
            started.set()
            release.wait(2)
            return report

        with patch.object(queries, "query_report", side_effect=wait_query):
            try:
                self.page.query()
                self.wait_for(started.is_set)
                self.page._partial_check.setChecked(False)
            finally:
                release.set()
            self.wait_for(lambda: not self.page._query_busy)
        self.assertIsNone(self.page._report)
        self.assertFalse(self.page._excel_button.isEnabled())
        self.assertIn("本次结果未显示", self.page._snapshot_label.text())

    def test_query_error_does_not_allow_export_of_old_data(self):
        self.show_report()
        with patch.object(queries, "query_report", side_effect=ValueError("读取失败")):
            self.page.query()
            self.wait_for(lambda: not self.page._query_busy)
        self.assertFalse(self.page._excel_button.isEnabled())
        self.assertIn("查询失败", self.page._snapshot_label.text())

    def test_invalid_date_range_never_starts_query(self):
        self.page._start_date.setDate(date(2026, 10, 20))
        with (
            patch.object(queries, "query_report") as query,
            patch("gui.app.pages.auxiliary_queries_page.msgbox.warning"),
        ):
            self.page.query()
        query.assert_not_called()

    def test_default_date_is_today_and_quick_ranges_cover_yesterday(self):
        self.assertEqual(self.page._criteria()[:2], (NOW.date(), NOW.date()))
        cases = [
            ("yesterday", date(2026, 10, 13), date(2026, 10, 13)),
            ("today", date(2026, 10, 14), date(2026, 10, 14)),
            ("week", date(2026, 10, 8), date(2026, 10, 14)),
            ("month", date(2026, 10, 1), date(2026, 10, 14)),
        ]
        with patch.object(self.page, "query") as query:
            for mode, start, end in cases:
                with self.subTest(mode=mode):
                    self.page._quick_range(mode)
                    self.assertEqual(self.page._criteria()[:2], (start, end))
            self.assertEqual(query.call_count, 4)

    def test_yesterday_button_reads_existing_orders_and_updates_all_tables(self):
        self.insert("orders", order_no="old-yesterday", submitted_at="2026-10-13T00:00:00Z")
        button = next(b for b in self.page.findChildren(QPushButton) if b.text() == "昨天")
        with patch.object(queries.db, "DB_PATH", self.path):
            button.click()
            self.wait_for(lambda: self.page._report is not None)
        self.assertEqual(self.page._criteria()[:2], (date(2026, 10, 13), date(2026, 10, 13)))
        report = self.page._report
        self.assertEqual(sum(row["total"] for row in report["daily"]), 1)
        self.assertEqual([row["order_no"] for row in report["details"]], ["old-yesterday"])
        self.assertEqual(sum(row["submitted"] for row in report["summary"]), 1)
        self.assertEqual(self.page._tabs.tabText(0), "提交数量（2026-10-13）")
        self.assertIn("2026-10-13", self.page._count_label.text())
        self.assertNotIn("今日提交", self.page._count_label.text())

    def test_custom_historical_date_queries_saved_submission_time(self):
        self.insert("aftersale_appeal_records", order_no="old-history", record_key="old",
                    status="申诉成功", platform_appealed_at="2025-09-02 12:30:00")
        self.show_report()
        self.page._start_date.setDate(date(2025, 9, 2))
        self.page._end_date.setDate(date(2025, 9, 2))
        self.assertFalse(self.page._csv_button.isEnabled())
        with patch.object(queries.db, "DB_PATH", self.path):
            self.page.query()
            self.wait_for(lambda: not self.page._query_busy)
        self.assertEqual(sum(row["total"] for row in self.page._report["daily"]), 1)
        self.assertEqual(self.page._report["details"][0]["order_no"], "old-history")
        self.assertTrue(self.page._excel_button.isEnabled())
        self.assertEqual(self.page._tabs.tabText(0), "提交数量（2025-09-02）")

    def test_numeric_sort_and_percent_display(self):
        model = QueryTableModel([("pass_rate", "通过率")])
        model.set_rows([{"pass_rate": 0.1}, {"pass_rate": None}, {"pass_rate": 0.9}])
        model.sort(0, Qt.DescendingOrder)
        self.assertEqual([row["pass_rate"] for row in model.rows], [0.9, 0.1, None])
        self.assertEqual(model.data(model.index(0, 0)), "90.0%")
        self.assertEqual(model.data(model.index(2, 0)), "—")


class _FakePage(QWidget):
    def __init__(self):
        super().__init__()
        self._login_check = None
        self.exporting = False
        self.active = False

    def _batch_is_running(self):
        return False

    def is_running(self):
        return False

    def is_exporting_captcha_diagnostics(self):
        return False

    def is_exporting(self):
        return self.exporting

    def set_active(self, active):
        self.active = active

    def set_polling_active(self, active):
        self.active = active


class AuxiliaryNavigationTests(_DatabaseFixture):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_navigation_is_below_reports_and_export_blocks_teardown(self):
        with (
            patch.object(main_window, "DashboardPage", _FakePage),
            patch.object(main_window, "ShopsPage", _FakePage),
            patch.object(main_window, "RunPage", _FakePage),
            patch.object(main_window, "ReportsPage", _FakePage),
            patch.object(main_window, "AuxiliaryQueriesPage", _FakePage),
        ):
            window = main_window.MainWindow()
        self.addCleanup(window.close)
        self.assertEqual(window._nav.item(4).text(), "辅助查询功能")
        self.assertIn("申诉", window._nav.item(3).text())
        self.assertEqual(window._stack.count(), 5)
        window._nav.setCurrentRow(4)
        self.assertTrue(window._auxiliary_page.active)
        self.assertFalse(window._reports_page.active)
        window._auxiliary_page.exporting = True
        event = MagicMock()
        window.closeEvent(event)
        event.ignore.assert_called_once()
        window._auxiliary_page.exporting = False


if __name__ == "__main__":
    import unittest

    unittest.main()
