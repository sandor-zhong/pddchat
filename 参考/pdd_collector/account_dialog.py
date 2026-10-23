"""Editable account staging grid; only encrypted credentials are committed to the database."""
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QKeySequence
from PyQt5.QtWidgets import (QApplication, QComboBox, QDialog, QDialogButtonBox, QFileDialog,
    QFormLayout, QHBoxLayout, QHeaderView, QLabel, QLineEdit, QMessageBox, QPushButton,
    QShortcut, QSpinBox, QTableWidget, QTableWidgetItem, QVBoxLayout)

from .account_import import normalize_accounts, rows_from_text, read_excel_rows
from .password_widgets import eye_icon, PasswordItem, PasswordDelegate


class ExcelColumnsDialog(QDialog):
    def __init__(self, path, parent=None):
        super().__init__(parent)
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        from .account_import import ALIASES
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            self.sheets = {}
            for sheet in workbook.worksheets:
                header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
                mapping = {field: next((i for i, value in enumerate(header)
                    if str(value or "").strip().lower().rstrip(":：") in aliases), None)
                    for field, aliases in ALIASES.items()}
                self.sheets[sheet.title] = (sheet.max_column, mapping)
        finally:
            workbook.close()
        self.setWindowTitle("选择 Excel 列")
        self.path = path
        self.get_column_letter = get_column_letter
        layout = QFormLayout(self)
        self.sheet = QComboBox()
        self.sheet.addItems(list(self.sheets))
        layout.addRow("工作表", self.sheet)
        self.start_row = QSpinBox()
        self.start_row.setRange(1, 1048576)
        self.start_row.setValue(2)
        layout.addRow("首行数据行号", self.start_row)
        self.columns = {}
        for field, label in (("username", "账号列"), ("password", "密码列"),
                             ("display_name", "店铺名列"), ("group_name", "分组列")):
            combo = QComboBox()
            self.columns[field] = combo
            layout.addRow(label, combo)
        self.sheet.currentTextChanged.connect(self.change_sheet)
        self.change_sheet()
        hint = QLabel("店铺名可不选择：新增店铺使用账号，已有店铺保留原名。\n账号和密码均为空的行自动跳过；原文件不作修改。")
        hint.setWordWrap(True)
        layout.addRow(hint)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("读取到列表")
        buttons.button(QDialogButtonBox.Cancel).setText("取消")
        buttons.accepted.connect(self.validate)
        buttons.rejected.connect(self.reject)
        layout.addRow(buttons)
        self.rows = []

    def change_sheet(self):
        count, mapping = self.sheets[self.sheet.currentText()]
        for field, combo in self.columns.items():
            combo.clear()
            if field in ("display_name", "group_name"):
                combo.addItem("不填写（新增店铺使用账号）" if field == "display_name" else "不设置", None)
            for index in range(count):
                combo.addItem(f"{self.get_column_letter(index + 1)} 列", index)
            default = mapping[field]
            if default is None and field in ("username", "password"):
                default = 0 if field == "username" else min(1, count - 1)
            combo.setCurrentIndex(max(0, combo.findData(default)))

    def validate(self):
        columns = {field: combo.currentData() for field, combo in self.columns.items()}
        chosen = [value for value in columns.values() if value is not None]
        if len(chosen) != len(set(chosen)):
            QMessageBox.information(self, "请检查列", "账号、密码、店铺名和分组需要选择不同的列。")
            return
        try:
            self.rows = read_excel_rows(self.path, self.sheet.currentText(), self.start_row.value(), columns)
        except Exception as exc:
            QMessageBox.warning(self, "读取失败", str(exc) if isinstance(exc, ValueError) else "无法读取文件，请检查格式与列设置。")
            return
        self.accept()


class AccountsDialog(QDialog):
    def __init__(self, db, parent=None):
        super().__init__(parent)
        self.db, self.keys, self.login_after_save = db, [], False
        self.setWindowTitle("账号列表录入 / 批量粘贴")
        self.resize(920, 600)
        layout = QVBoxLayout(self)
        hint = QLabel("每行一家：账号、密码必填；第 3 列店铺名、第 4 列分组可选。支持直接输入或粘贴 Excel 两至四列（Ctrl+V）。\n按登录账号去重，同批重复账号以最后一行为准；已有账号更新原记录，名称留空保留原名。点击眼睛切换密码显示。")
        hint.setWordWrap(True)
        layout.addWidget(hint)
        toolbar = QHBoxLayout()
        for label, slot in (("新增一行", self.add_row), ("粘贴多行（追加）", self.paste),
                            ("读取 Excel…", self.load_excel), ("移除选中行", self.remove_rows)):
            button = QPushButton(label)
            button.clicked.connect(slot)
            toolbar.addWidget(button)
        toolbar.addStretch()
        layout.addLayout(toolbar)
        self.table = QTableWidget(0, 4)
        self.table.passwords_visible = getattr(parent, 'passwords_visible', True)
        self.table.setItemDelegateForColumn(1, PasswordDelegate(self.table))
        self.password_toggle = QPushButton()
        self.password_toggle.clicked.connect(self.toggle_passwords)
        toolbar.addWidget(self.password_toggle)
        self.update_password_button()
        self.table.setHorizontalHeaderLabels(["账号", "密码", "店铺名（可选）", "分组（可选）"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setAlternatingRowColors(True)
        layout.addWidget(self.table, 1)
        self.shortcut = QShortcut(QKeySequence.Paste, self.table)
        self.shortcut.setContext(Qt.WidgetWithChildrenShortcut)
        self.shortcut.activated.connect(self.paste)
        self.count = QLabel()
        self.table.itemChanged.connect(self.update_count)
        layout.addWidget(self.count)
        buttons = QHBoxLayout()
        buttons.addStretch()
        for label, slot in (("保存列表", lambda: self.save(False)),
                            ("保存并批量登录", lambda: self.save(True)), ("取消", self.reject)):
            button = QPushButton(label)
            button.clicked.connect(slot)
            buttons.addWidget(button)
        layout.addLayout(buttons)
        self.add_row()

    def add_row(self, values=None):
        values = values if isinstance(values, (list, tuple)) else []
        row = self.table.rowCount()
        self.table.insertRow(row)
        for col in range(4):
            item = PasswordItem() if col == 1 else QTableWidgetItem()
            item.setData(Qt.EditRole, str(values[col]) if col < len(values) and values[col] is not None else "")
            self.table.setItem(row, col, item)
        self.update_count()

    def rows(self):
        return [[self.table.item(row, col).data(Qt.EditRole) if self.table.item(row, col) else ""
                 for col in range(4)] for row in range(self.table.rowCount())]

    def update_password_button(self):
        show = self.table.passwords_visible
        self.password_toggle.setIcon(eye_icon(show))
        self.password_toggle.setText('隐藏密码' if show else '显示密码')
        self.password_toggle.setToolTip(self.password_toggle.text())

    def toggle_passwords(self):
        self.table.setFocus()
        self.table.passwords_visible = not self.table.passwords_visible
        self.update_password_button()
        self.table.viewport().update()

    def update_count(self):
        self.count.setText(f"待保存 {sum(any(cell for cell in row) for row in self.rows())} 行 · 双击可修改 · 店铺名和分组可留空")

    def append_rows(self, rows):
        if any(len(row) > 4 for row in rows):
            raise ValueError("最多粘贴 4 列：账号、密码、店铺名、分组。请先在 Excel 选择对应列。")
        current_rows = self.rows()
        for row in reversed(range(self.table.rowCount())):
            if not any(current_rows[row]):
                self.table.removeRow(row)
        for row in rows:
            self.add_row(row)
        self.update_count()

    def paste(self):
        text = QApplication.clipboard().text()
        # A single value is pasted into the current cell; structured rows are appended.
        editor = QApplication.focusWidget()
        if isinstance(editor, QLineEdit) and "\t" not in text and "\n" not in text:
            editor.paste()
            return
        if "\t" not in text and "\n" not in text and "," not in text and self.table.currentRow() >= 0:
            editor = QApplication.focusWidget()
            if isinstance(editor, QLineEdit):
                editor.paste()
            else:
                self.table.item(self.table.currentRow(), self.table.currentColumn()).setData(Qt.EditRole, text)
            return
        try:
            self.append_rows(rows_from_text(text))
        except Exception as exc:
            QMessageBox.information(self, "粘贴失败", str(exc) if isinstance(exc, ValueError) else "无法识别粘贴内容，请使用 Excel 两至四列或 CSV。")

    def load_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "读取账号列表", "", "Excel 工作簿 (*.xlsx)")
        if path:
            try:
                dialog = ExcelColumnsDialog(path, self)
                if dialog.exec_() == QDialog.Accepted:
                    self.append_rows(dialog.rows)
            except Exception:
                QMessageBox.warning(self, "读取失败", "无法打开文件，请选择未加密的 .xlsx 工作簿。")

    def remove_rows(self):
        for row in sorted({index.row() for index in self.table.selectedIndexes()}, reverse=True):
            self.table.removeRow(row)
        if not self.table.rowCount():
            self.add_row()
        self.update_count()

    def save(self, login):
        self.table.setFocus()  # Commit any active cell editor before reading values.
        try:
            records = normalize_accounts(self.rows())
            self.keys = self.db.import_accounts(records)
        except Exception as exc:
            QMessageBox.warning(self, "保存未完成", str(exc))
            return
        self.login_after_save = login
        self.accept()
