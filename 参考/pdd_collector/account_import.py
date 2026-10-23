"""Account input normalization. Passwords are never trimmed or included in errors."""
import csv
from hashlib import sha256
import io


ALIASES = {
    "username": {"账号", "帐号", "账户", "登录账号", "账号名", "子账号", "username", "account"},
    "password": {"密码", "登录密码", "账号密码", "password"},
    "display_name": {"店铺", "店铺名", "店铺名称", "店名", "display_name", "shop_name"},
    "group_name": {"分组", "所属组", "店铺分组", "group_name"},
}


def header_mapping(row):
    normalized = [str(value or "").strip().lower().rstrip(":：").removesuffix("（可选）").removesuffix("(可选)") for value in row]
    mapping = {field: next((i for i, value in enumerate(normalized) if value in names), None)
               for field, names in ALIASES.items()}
    return mapping if all(mapping[key] is not None for key in ("username", "password")) else None

def normalize_accounts(rows):
    result = {}
    for number, row in enumerate(rows, 1):
        if not any(value is not None and str(value) != "" for value in row):
            continue
        if len(row) < 2 or len(row) > 4:
            raise ValueError(f"第 {number} 行需要账号、密码，可选第 3 列店铺名、第 4 列分组。")
        username = str(row[0] or "").strip()
        password = "" if row[1] is None else str(row[1])
        name = str(row[2] or "").strip() if len(row) > 2 else ""
        if not username or not password:
            raise ValueError(f"第 {number} 行的账号和密码均不能为空。")
        if len(username) > 200 or len(name) > 120:
            raise ValueError(f"第 {number} 行账号或店铺名过长。")
        # Preserve first occurrence order; the last row supplies this account's values.
        result[username] = dict(username=username, password=password, display_name=name,
                                group_name=str(row[3] or "").strip() if len(row) > 3 else "")
    if not result:
        raise ValueError("没有可保存的账号，请输入或粘贴账号和密码，店铺名可选。")
    return list(result.values())

def rows_from_text(text):
    text = text.lstrip("\ufeff")
    rows = list(csv.reader(io.StringIO(text), delimiter="\t" if "\t" in text else ","))
    rows = [row for row in rows if any(cell != "" for cell in row)]
    if not rows:
        return []
    mapping = header_mapping(rows[0])
    if mapping:
        return [[row[mapping[key]] if mapping[key] is not None and mapping[key] < len(row) else ""
                 for key in ("username", "password", "display_name", "group_name")] for row in rows[1:]]
    return rows


def account_key(username):
    return "account-" + sha256(username.encode("utf-8")).hexdigest()[:24]


def read_excel_rows(path, sheet_name, start_row, columns):
    """Read selected columns without interpreting unrelated notes or dates as credentials."""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    rows = []
    try:
        sheet = workbook[sheet_name]
        for number, cells in enumerate(sheet.iter_rows(min_row=start_row), start_row):
            values = {}
            for field, index in columns.items():
                cell = cells[index] if index is not None and index < len(cells) else None
                value = cell.value if cell is not None else None
                if cell is not None and cell.data_type == "f":
                    raise ValueError(f"第 {number} 行选中列含公式，请转换为文本值后读取。")
                if field in ("username", "password") and value is not None and not isinstance(value, str):
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value):
                        text = str(int(value))
                        if len(text.lstrip("-")) > 15:
                            raise ValueError(f"第 {number} 行数字账号或密码过长，请改为文本并核对精度。")
                        fmt = cell.number_format
                        value = text.zfill(len(fmt)) if fmt and set(fmt) == {"0"} else text
                    else:
                        raise ValueError(f"第 {number} 行账号或密码不是文本或整数，请检查列选择。")
                values[field] = "" if value is None else str(value)
            if not values["username"] and not values["password"]:
                continue
            rows.append([values[field] for field in ("username", "password", "display_name", "group_name")])
    finally:
        workbook.close()
    if not rows:
        raise ValueError("所选范围没有账号，请检查工作表、数据起始行和列设置。")
    return rows


def read_excel_accounts(path):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        raise ValueError("无法打开 Excel，请选择未加密的 .xlsx 文件。") from None
    records = []
    try:
        for sheet in workbook.worksheets:
            mapping = None
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                if mapping is None:
                    mapping = header_mapping([cell.value for cell in cells])
                    if mapping is not None:
                        continue
                    if row_number >= 20:
                        break
                    continue
                values = []
                for field in ("username", "password", "display_name", "group_name"):
                    index = mapping[field]
                    cell = cells[index] if index is not None and index < len(cells) else None
                    value = cell.value if cell is not None else None
                    if cell is not None and cell.data_type == "f":
                        raise ValueError(f"工作表 {sheet.title} 第 {row_number} 行包含公式，请转换为文本值后导入。")
                    if field in ("username", "password") and value is not None and not isinstance(value, str):
                        # Reconstruct an explicit zero-padding number format; refuse lossy identifiers.
                        if isinstance(value, (int, float)) and not isinstance(value, bool) and int(value) == value:
                            text = str(int(value))
                            if len(text.lstrip('-')) > 15:
                                raise ValueError(f"工作表 {sheet.title} 第 {row_number} 行的数字账号/密码过长，请改为文本核对精度。")
                            fmt = cell.number_format
                            value = text.zfill(len(fmt)) if fmt and set(fmt) == {"0"} else text
                        else:
                            raise ValueError(f"工作表 {sheet.title} 第 {row_number} 行的账号/密码格式不支持，请改为文本。")
                    values.append(value)
                if any(value is not None and str(value) != "" for value in values):
                    records.append(values)
    finally:
        workbook.close()
    return normalize_accounts(records)
