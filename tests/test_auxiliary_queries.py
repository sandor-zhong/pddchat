from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

import db
from gui.app.services import auxiliary_queries as queries

NOW = datetime(2026, 10, 14, 12, tzinfo=queries.BEIJING)


class _DatabaseFixture(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        self.path = self.root / "state/pddchat.db"
        self.path.parent.mkdir()
        with closing(sqlite3.connect(self.path)) as connection:
            db._init_schema(connection)

    def insert(self, table, **values):
        defaults = {
            "shop": "店铺A",
            "order_no": "261001-123456789",
            "status": "approved",
            "updated_at": "2026-10-14T00:00:00+00:00",
        }
        if table != "orders":
            defaults.update(
                record_key=values.get("record_key", "r1"),
                discovered_at="2026-10-01T00:00:00+00:00",
            )
        if table == "reconsideration_records":
            defaults["last_seen_at"] = defaults["updated_at"]
        if table == "appeal_records":
            defaults.pop("discovered_at", None)
        defaults.update(values)
        with closing(sqlite3.connect(self.path)) as connection:
            names = ",".join(defaults)
            placeholders = ",".join("?" for _value in defaults)
            connection.execute(
                f"INSERT INTO {table} ({names}) VALUES ({placeholders})",
                list(defaults.values()),
            )
            connection.commit()

    def report(self, **kwargs):
        return queries.query_report(
            date(2026, 10, 1), date(2026, 10, 14), now=NOW, database=self.path, **kwargs
        )


class AuxiliaryQueryTests(_DatabaseFixture):
    def test_three_types_counts_rates_and_simulated_exclusion(self):
        for order, status in [
            ("a", "approved"),
            ("b", "rejected"),
            ("c", "submitted"),
            ("d", "simulated"),
        ]:
            self.insert(
                "orders",
                order_no=order,
                status=status,
                submitted_at="2026-10-14T00:00:00+00:00",
            )
        for key, status in [("r1", "申诉成功"), ("r2", "申诉失败"), ("r3", "模拟成功")]:
            self.insert(
                "reconsideration_records",
                record_key=key,
                status=status,
                submitted_at="2026-10-14T01:00:00+00:00",
            )
        for key, status in [
            ("s1", "申诉成功"),
            ("s2", "部分通过"),
            ("s3", "申诉失败"),
            ("s4", "审核中"),
            ("s5", "已撤回"),
            ("s6", "材料已准备"),
        ]:
            self.insert(
                "aftersale_appeal_records",
                record_key=key,
                status=status,
                submitted_at="2026-10-14T02:00:00+00:00",
            )
        report = self.report()
        self.assertEqual(report["daily"][0]["total"], 10)
        self.assertEqual(
            [report["daily"][0][kind] for kind in queries.BUSINESS_TYPES], [3, 2, 5]
        )
        groups = {row["business_type"]: row for row in report["summary"]}
        self.assertEqual(groups["负向申诉"]["pass_rate"], 0.5)
        self.assertEqual(groups["复议"]["pass_rate"], 0.5)
        self.assertAlmostEqual(groups["售后申诉"]["pass_rate"], 2 / 3)
        self.assertEqual(groups["售后申诉"]["submission_pass_rate"], 2 / 5)
        self.assertEqual(groups["售后申诉"]["reviewed"], 3)
        self.assertEqual(groups["售后申诉"]["withdrawn"], 1)
        self.assertEqual(len(report["details"]), 10)
        strict = self.report(partial_pass=False)
        sale = next(
            row for row in strict["summary"] if row["business_type"] == "售后申诉"
        )
        self.assertEqual(sale["reviewed"], 3)
        self.assertAlmostEqual(sale["pass_rate"], 1 / 3)

    def test_beijing_start_end_boundaries_and_today(self):
        stamps = [
            "2026-09-30T15:59:59+00:00",
            "2026-09-30T16:00:00+00:00",
            "2026-10-01T15:59:59+00:00",
            "2026-10-01T16:00:00+00:00",
        ]
        for index, stamp in enumerate(stamps):
            self.insert("orders", order_no=str(index), submitted_at=stamp)
        report = queries.query_report(
            date(2026, 10, 1), date(2026, 10, 1), now=NOW, database=self.path
        )
        self.assertEqual({row["order_no"] for row in report["details"]}, {"1", "2"})
        self.assertEqual(report["details"][0]["submission_time"], "2026-10-01 23:59:59")
        for index, stamp in enumerate(
            [
                "2026-10-13T15:59:59Z",
                "2026-10-13T16:00:00Z",
                "2026-10-14T03:59:59Z",
                "2026-10-14T04:00:01Z",
            ]
        ):
            self.insert("orders", order_no="today-" + str(index), submitted_at=stamp)
        report = queries.query_report(
            NOW.date(), NOW.date(), now=NOW, database=self.path
        )
        self.assertEqual(report["daily"][0]["负向申诉"], 2)
        self.assertNotIn("today-3", {row["order_no"] for row in report["details"]})

    def test_platform_dates_complete_counts_and_local_date_has_priority(self):
        self.insert(
            "aftersale_appeal_records",
            status="申诉成功",
            platform_appealed_at="2026年10月14日 10:30",
        )
        self.insert(
            "reconsideration_records",
            record_key="r2",
            status="申诉成功",
            submitted_at="2026-09-01T00:00:00Z",
            platform_reconsidered_at="2026-10-14 10:00:00",
        )
        report = self.report()
        self.assertEqual(len(report["details"]), 1)
        self.assertEqual(report["details"][0]["time_source"], "平台提交时间")
        self.assertEqual(report["details"][0]["submission_time"], "2026-10-14 10:30:00")
        self.assertEqual(report["daily"][0]["total"], 1)

    def test_negative_order_dedup_but_distinct_reconsiderations_preserved(self):
        self.insert("orders", submitted_at="2026-10-14T00:00:00Z")
        for key in ("a", "b"):
            self.insert(
                "appeal_records",
                record_key=key,
                status="submitted",
                submitted_at="2026-10-14T00:00:00Z",
            )
            self.insert(
                "reconsideration_records",
                record_key=key,
                status="申诉成功",
                submitted_at="2026-10-14T00:00:00Z",
            )
        report = self.report()
        self.assertEqual(report["daily"][0]["负向申诉"], 1)
        self.assertEqual(report["daily"][0]["复议"], 2)
        self.assertEqual(len(report["details"]), 3)

    def test_child_record_submission_date_can_complete_negative_order_date(self):
        self.insert("orders", status="approved")
        self.insert(
            "appeal_records", status="submitted", submitted_at="2026-10-14T00:00:00Z"
        )
        report = self.report()
        self.assertEqual(report["daily"][0]["负向申诉"], 1)
        self.assertEqual(report["details"][0]["outcome"], "通过")

    def test_shop_and_type_filters_are_exact_and_parameterized(self):
        self.insert("orders", order_no="a", submitted_at="2026-10-14T00:00:00Z")
        name = "A' OR 1=1 --"
        self.insert(
            "orders", shop=name, order_no="b", submitted_at="2026-10-14T00:00:00Z"
        )
        self.insert(
            "reconsideration_records",
            shop=name,
            status="申诉成功",
            submitted_at="2026-10-14T00:00:00Z",
        )
        report = self.report(shop=name, business_type="复议")
        self.assertEqual(len(report["details"]), 1)
        self.assertEqual(report["details"][0]["shop"], name)
        self.assertEqual(report["daily"][0]["负向申诉"], 0)
        self.assertEqual(len(report["summary"]), 1)

    def test_zero_submission_shops_are_listed_with_no_fake_zero_percent(self):
        profiles = self.root / "state/profiles/空店铺"
        profiles.mkdir(parents=True)
        imported = self.root / "gui/state/imported_shops.json"
        imported.parent.mkdir(parents=True)
        imported.write_text(json.dumps({"shops": {"导入店铺": {}}}), encoding="utf-8")
        report = self.report()
        self.assertEqual(
            {row["shop"] for row in report["daily"]}, {"空店铺", "导入店铺"}
        )
        self.assertTrue(all(row["total"] == 0 for row in report["daily"]))
        self.assertTrue(all(row["pass_rate"] is None for row in report["summary"]))

    def test_missing_and_invalid_dates_are_reported_not_replaced_with_update_time(self):
        self.insert("reconsideration_records", record_key="a", status="申诉成功")
        self.insert(
            "aftersale_appeal_records",
            record_key="b",
            status="申诉成功",
            submitted_at="bad date",
        )
        report = self.report()
        self.assertEqual(report["missing_time"], 2)
        self.assertFalse(report["details"])
        self.assertEqual(report["daily"][0]["total"], 0)

    def test_today_yesterday_and_arbitrary_day_count_existing_orders_by_submission(self):
        for day in (12, 13, 14):
            for table, status in (("orders", "approved"), ("reconsideration_records", "申诉失败"),
                                  ("aftersale_appeal_records", "部分通过")):
                self.insert(table, status=status, order_no=f"{table}-{day}",
                            submitted_at=f"2026-10-{day:02}T10:00:00+08:00",
                            updated_at="2026-10-14T00:00:00Z")
        for day in (12, 13, 14):
            with self.subTest(day=day):
                selected = date(2026, 10, day)
                report = queries.query_report(selected, selected, now=NOW, database=self.path)
                self.assertEqual([report["daily"][0][kind] for kind in queries.BUSINESS_TYPES], [1, 1, 1])
                self.assertEqual(report["daily"][0]["total"], 3)
                self.assertEqual(len(report["details"]), 3)
                self.assertEqual(sum(row["submitted"] for row in report["summary"]), 3)
                self.assertTrue(all(row["submission_time"].startswith(selected.isoformat()) for row in report["details"]))

    def test_range_counts_and_details_share_submission_period(self):
        for day in (12, 13, 14):
            self.insert("orders", order_no=str(day), submitted_at=f"2026-10-{day:02}T10:00:00+08:00")
        self.insert("reconsideration_records", status="申诉成功", platform_reconsidered_at="2026-10-13 23:59:59")
        self.insert("aftersale_appeal_records", status="申诉成功", platform_appealed_at="2026-10-14 00:00:00")
        report = queries.query_report(date(2026, 10, 12), date(2026, 10, 13), now=NOW, database=self.path)
        self.assertEqual([report["daily"][0][kind] for kind in queries.BUSINESS_TYPES], [2, 1, 0])
        self.assertEqual(sum(row["total"] for row in report["daily"]), len(report["details"]))
        self.assertEqual(len(report["details"]), 3)

    def test_old_platform_and_child_record_dates_are_queryable(self):
        self.insert("orders", status="approved", updated_at="2026-10-14T00:00:00Z")
        self.insert("appeal_records", status="submitted", submitted_at="2025-09-01T16:00:00Z")
        self.insert("reconsideration_records", status="申诉成功", platform_reconsidered_at="2025年09月02日 10:30")
        self.insert("aftersale_appeal_records", status="申诉成功", platform_appealed_at="2025-09-02 23:59:59")
        report = queries.query_report(date(2025, 9, 2), date(2025, 9, 2), now=NOW, database=self.path)
        self.assertEqual([report["daily"][0][kind] for kind in queries.BUSINESS_TYPES], [1, 1, 1])
        self.assertEqual(len(report["details"]), 3)
        self.assertEqual(sum(row["submitted"] for row in report["summary"]), 3)
        self.assertEqual({row["time_source"] for row in report["details"]}, {"程序提交时间", "平台提交时间"})

    def test_read_only_query_does_not_modify_database(self):
        self.insert("orders", submitted_at="2026-10-14T00:00:00Z")
        before = hashlib.sha256(self.path.read_bytes()).digest()
        self.report()
        self.assertEqual(before, hashlib.sha256(self.path.read_bytes()).digest())
        with self.assertRaises(ValueError):
            queries.query_report(
                date(2026, 10, 15), date(2026, 10, 1), database=self.path
            )

    def test_each_report_reads_one_consistent_database_snapshot(self):
        with closing(sqlite3.connect(self.path)) as connection:
            connection.execute("PRAGMA journal_mode=WAL")
        self.insert("orders", order_no="initial", submitted_at="2026-10-14T00:00:00Z")
        original = queries._known_shops

        def after_snapshot(connection, root):
            names = original(connection, root)
            self.insert(
                "aftersale_appeal_records",
                status="申诉成功",
                submitted_at="2026-10-14T00:00:00Z",
            )
            return names

        with patch.object(queries, "_known_shops", side_effect=after_snapshot):
            report = self.report()
        self.assertEqual(len(report["details"]), 1)
        self.assertEqual(len(self.report()["details"]), 2)


class AuxiliaryExportTests(_DatabaseFixture):
    def make_report(self):
        self.insert(
            "aftersale_appeal_records",
            status="部分通过",
            submitted_at="2026-10-14T00:00:00Z",
            order_no="000012345678901234567890",
            goods_title="=1+1",
        )
        return self.report()

    def test_workbook_contains_matching_summary_detail_and_literal_ids(self):
        from openpyxl import load_workbook

        from gui.app.services import auxiliary_export as export

        report = self.make_report()
        target = self.root / "report.xlsx"
        export.export_report(report, target)
        workbook = load_workbook(target)
        self.addCleanup(workbook.close)
        self.assertEqual(
            workbook.sheetnames,
            ["查询说明", "提交数量", "通过率汇总", "订单结果明细"],
        )
        sheet = workbook["订单结果明细"]
        columns = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.max_row - 1, len(report["details"]))
        self.assertEqual(
            sheet.cell(2, columns["订单号"]).value, "000012345678901234567890"
        )
        self.assertEqual(sheet.cell(2, columns["商品名称"]).value, "=1+1")
        self.assertEqual(sheet.cell(2, columns["商品名称"]).data_type, "s")
        self.assertEqual(
            sheet.cell(2, columns["统计提交时间（北京时间）"]).value,
            "2026-10-14 08:00:00",
        )
        self.assertEqual(workbook["通过率汇总"].max_row - 1, len(report["summary"]))
        metadata = dict(workbook["查询说明"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(metadata["明细记录数"], 1)
        self.assertEqual(metadata["查询开始日期（北京时间）"], "2026-10-01")

    def test_historical_day_export_labels_and_orders_match_selected_date(self):
        from openpyxl import load_workbook
        from gui.app.services import auxiliary_export as export
        self.insert("orders", order_no="old-order", submitted_at="2025-01-01T16:00:00Z")
        report = queries.query_report(date(2025, 1, 2), date(2025, 1, 2), now=NOW, database=self.path)
        target = self.root / "historical.xlsx"
        export.export_report(report, target)
        with target.open("rb") as handle:
            workbook = load_workbook(handle)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook["提交数量"].cell(2, 2).value, 1)
            self.assertEqual(workbook["订单结果明细"].max_row - 1, 1)
            metadata = dict(workbook["查询说明"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(metadata["提交数量统计范围"], "2025-01-02 至 2025-01-02（包含结束当天）")
            self.assertEqual(metadata["查询开始日期（北京时间）"], "2025-01-02")

    def test_csv_has_bom_dates_every_order_and_safe_literal_text(self):
        from gui.app.services import auxiliary_export as export

        report = self.make_report()
        target = self.root / "report.csv"
        export.export_report(report, target, "csv")
        self.assertTrue(target.read_bytes().startswith(b"\xef\xbb\xbf"))
        with target.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(len(rows), len(report["details"]))
        self.assertEqual(rows[0]["订单号"], "000012345678901234567890")
        self.assertEqual(rows[0]["商品名称"], "'=1+1")
        self.assertEqual(rows[0]["查询结束日期"], "2026-10-14")

    def test_failed_export_preserves_existing_file_and_removes_temporary(self):
        from gui.app.services import auxiliary_export as export

        report = self.make_report()
        target = self.root / "existing.xlsx"
        target.write_bytes(b"original")
        with patch.object(
            export.os, "replace", side_effect=PermissionError("file is open")
        ):
            with self.assertRaises(PermissionError):
                export.export_report(report, target)
        self.assertEqual(target.read_bytes(), b"original")
        self.assertFalse(list(self.root.glob(".appeal-export-*")))


if __name__ == "__main__":
    unittest.main()
